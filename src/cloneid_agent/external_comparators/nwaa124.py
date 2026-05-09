"""Extraction utilities for the Li et al. NSR 2021 nwaa124 supplement."""

from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path
import shutil
import zipfile
from typing import Any
from xml.etree import ElementTree as ET
from xml.sax.saxutils import escape


RECORD_STATUS_LABELS = (
    "structured_numeric_table",
    "model_formula_in_caption_or_methods",
    "embedded_plot_or_representative_image",
    "method_text_only",
    "not_available_in_archive",
    "not_event_linked",
    "not_agent_ready_without_manual_reconstruction",
)

DOCX_EXTRACTION_ENGINE = "python-docx_optional_openxml_fallback"

WORD_NS = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}


@dataclass(frozen=True)
class ArchiveSource:
    """Resolved external supplement source."""

    requested_path: str
    resolved_path: Path
    source_type: str
    used_fallback: bool


def _candidate_fallback_paths(requested: Path) -> list[Path]:
    cwd = Path.cwd()
    candidates = [
        cwd / "data" / "nwaa124_supplement_file.zip",
        cwd / "data" / "nwaa124_supplement_file",
        cwd.parent / "cloneid_physicell_agent_codex" / "data" / "nwaa124_supplement_file.zip",
        cwd.parent / "cloneid_physicell_agent_codex" / "data" / "nwaa124_supplement_file",
        Path.home() / "Documents" / "GitHub" / "cloneid_physicell_agent_codex" / "data" / "nwaa124_supplement_file.zip",
        Path.home() / "Documents" / "GitHub" / "cloneid_physicell_agent_codex" / "data" / "nwaa124_supplement_file",
    ]
    if requested.name:
        candidates.extend(
            [
                cwd / requested.name,
                cwd / "data" / requested.name,
                cwd.parent / "cloneid_physicell_agent_codex" / "data" / requested.name,
            ]
        )
    return candidates


def resolve_external_archive_path(external_zip_or_dir: str | Path | None) -> ArchiveSource:
    """Resolve the requested NSR supplement zip or directory.

    The benchmark accepts a zip file or an already-unpacked supplement directory.
    In local Codex runs the uploaded `/mnt/data` path may not be mounted, so the
    known repository data directory is used as a deterministic fallback.
    """

    requested = Path(external_zip_or_dir or "")
    if external_zip_or_dir and requested.exists():
        source_type = "zip" if requested.suffix.lower() == ".zip" else "directory"
        return ArchiveSource(str(external_zip_or_dir), requested, source_type, False)

    for candidate in _candidate_fallback_paths(requested):
        if candidate.exists():
            source_type = "zip" if candidate.suffix.lower() == ".zip" else "directory"
            return ArchiveSource(str(external_zip_or_dir or ""), candidate, source_type, True)

    raise FileNotFoundError(
        "Could not locate the nwaa124 supplement archive or directory. "
        f"Requested: {external_zip_or_dir!r}"
    )


def classify_archive_member(path: str | Path) -> str:
    suffix = Path(path).suffix.lower()
    if suffix == ".docx":
        return "docx"
    if suffix == ".txt":
        return "txt"
    if suffix == ".vcf":
        return "vcf"
    if suffix == ".zip":
        return "zip"
    return "other"


def _inventory_zip(path: Path) -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []
    with zipfile.ZipFile(path) as archive:
        for info in archive.infolist():
            if info.is_dir():
                continue
            rows.append(
                {
                    "path": info.filename,
                    "name": Path(info.filename).name,
                    "size_bytes": info.file_size,
                    "kind": classify_archive_member(info.filename),
                }
            )
    return sorted(rows, key=lambda row: row["path"])


def _inventory_directory(path: Path) -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []
    for member in sorted(path.rglob("*")):
        if not member.is_file():
            continue
        rel = member.relative_to(path).as_posix()
        rows.append(
            {
                "path": rel,
                "name": member.name,
                "size_bytes": member.stat().st_size,
                "kind": classify_archive_member(member),
            }
        )
    return rows


def inventory_external_archive(external_zip_or_dir: str | Path | None) -> dict[str, Any]:
    """List all external archive files and classify docx/txt/vcf members."""

    source = resolve_external_archive_path(external_zip_or_dir)
    files = _inventory_zip(source.resolved_path) if source.source_type == "zip" else _inventory_directory(source.resolved_path)
    counts: dict[str, int] = {}
    for row in files:
        counts[row["kind"]] = counts.get(row["kind"], 0) + 1
    return {
        "comparator": "Li et al. National Science Review 2021 nwaa124",
        "requested_path": source.requested_path,
        "resolved_path": str(source.resolved_path),
        "source_type": source.source_type,
        "used_fallback": source.used_fallback,
        "file_count": len(files),
        "kind_counts": counts,
        "files": files,
    }


def materialize_archive(external_zip_or_dir: str | Path | None, work_dir: str | Path) -> Path:
    """Return a directory containing archive members."""

    source = resolve_external_archive_path(external_zip_or_dir)
    if source.source_type == "directory":
        return source.resolved_path

    target = Path(work_dir) / "nwaa124_unpacked"
    target.mkdir(parents=True, exist_ok=True)
    with zipfile.ZipFile(source.resolved_path) as archive:
        archive.extractall(target)
    return target


def _paragraph_text(paragraph: ET.Element) -> str:
    return "".join(node.text or "" for node in paragraph.findall(".//w:t", WORD_NS)).strip()


def _table_rows(table: ET.Element) -> list[list[str]]:
    rows: list[list[str]] = []
    for table_row in table.findall(".//w:tr", WORD_NS):
        cells: list[str] = []
        for table_cell in table_row.findall("./w:tc", WORD_NS):
            cell_paragraphs = []
            for paragraph in table_cell.findall(".//w:p", WORD_NS):
                text = _paragraph_text(paragraph)
                if text:
                    cell_paragraphs.append(text)
            cells.append(" ".join(cell_paragraphs).strip())
        rows.append(cells)
    return rows


def _extract_openxml_docx(docx_path: Path) -> dict[str, Any]:
    with zipfile.ZipFile(docx_path) as archive:
        document_xml = archive.read("word/document.xml")
    root = ET.fromstring(document_xml)
    body = root.find("w:body", WORD_NS)
    paragraphs: list[str] = []
    tables: list[dict[str, Any]] = []
    document_order: list[dict[str, Any]] = []
    if body is None:
        return {"paragraphs": paragraphs, "tables": tables, "document_order": document_order}

    table_index = 0
    for child in body:
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "p":
            text = _paragraph_text(child)
            if text:
                paragraphs.append(text)
                document_order.append({"type": "paragraph", "text": text})
        elif tag == "tbl":
            table_index += 1
            rows = _table_rows(child)
            tables.append({"table_index": table_index, "rows": rows})
            document_order.append({"type": "table", "table_index": table_index, "rows": rows})

    return {"paragraphs": paragraphs, "tables": tables, "document_order": document_order}


def _extract_with_python_docx(docx_path: Path) -> dict[str, Any] | None:
    try:
        import docx  # type: ignore[import-not-found]
    except Exception:
        return None

    try:
        document = docx.Document(str(docx_path))
    except Exception:
        return None

    paragraphs = [paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text.strip()]
    tables: list[dict[str, Any]] = []
    for idx, table in enumerate(document.tables, start=1):
        rows = [[cell.text.strip() for cell in row.cells] for row in table.rows]
        tables.append({"table_index": idx, "rows": rows})
    return {
        "paragraphs": paragraphs,
        "tables": tables,
        "document_order": [],
        "python_docx_available": True,
    }


def _extract_docx_media(docx_path: Path, media_output_dir: Path) -> list[dict[str, Any]]:
    media_rows: list[dict[str, Any]] = []
    media_output_dir.mkdir(parents=True, exist_ok=True)
    with zipfile.ZipFile(docx_path) as archive:
        for name in archive.namelist():
            if not name.startswith("word/media/"):
                continue
            member_name = Path(name).name
            target = media_output_dir / member_name
            with archive.open(name) as source, target.open("wb") as handle:
                shutil.copyfileobj(source, handle)
            media_rows.append(
                {
                    "archive_path": name,
                    "extracted_path": str(target),
                    "name": member_name,
                    "size_bytes": target.stat().st_size,
                    "record_status": "embedded_plot_or_representative_image",
                }
            )
    return sorted(media_rows, key=lambda row: row["archive_path"])


def extract_docx_payload(docx_path: str | Path, media_output_dir: str | Path) -> dict[str, Any]:
    """Extract docx paragraphs, tables, and embedded media."""

    path = Path(docx_path)
    openxml_payload = _extract_openxml_docx(path)
    python_docx_payload = _extract_with_python_docx(path)
    if python_docx_payload is not None:
        payload = python_docx_payload
        payload["document_order"] = openxml_payload["document_order"]
        engine = "python-docx+openxml_document_order"
    else:
        payload = openxml_payload
        payload["python_docx_available"] = False
        engine = "openxml_fallback"

    media_dir = Path(media_output_dir) / path.stem.replace(" ", "_")
    payload.update(
        {
            "docx_path": str(path),
            "docx_name": path.name,
            "extraction_engine": engine,
            "paragraph_count": len(payload["paragraphs"]),
            "table_count": len(payload["tables"]),
            "embedded_media": _extract_docx_media(path, media_dir),
        }
    )
    payload["embedded_media_count"] = len(payload["embedded_media"])
    return payload


def extract_all_docx_payloads(external_zip_or_dir: str | Path | None, output_dir: str | Path) -> list[dict[str, Any]]:
    """Materialize the archive and extract every docx member."""

    root = materialize_archive(external_zip_or_dir, Path(output_dir) / "_archive_work")
    docx_paths = sorted(root.rglob("*.docx"))
    media_output_dir = Path(output_dir) / "nwaa124_docx_media"
    return [extract_docx_payload(path, media_output_dir) for path in docx_paths]


def create_minimal_mock_supplement_fixture(output_dir: str | Path) -> Path:
    """Create a tiny labeled supplement-shaped fixture for missing-archive mock runs.

    This is not used when a real zip or directory is available. It exists so
    mock-mode integration tests can exercise the extraction path in dependency-
    constrained environments where `/mnt/data` is not mounted.
    """

    root = Path(output_dir) / "_mock_nwaa124_supplement_file"
    root.mkdir(parents=True, exist_ok=True)
    figure_paragraphs = [
        "Supplementary Figure 1 | Fitness of r and K cells compared with IN cells.",
        "Mixed populations were measured over passages.",
        "Supplementary Figure 4 | Observed dynamics of mixed populations initiated with 90% r cells and 10% K cells.",
        "The bars represent the proportion change of cell population by time.",
        "Supplementary Figure 5 | Predicted dynamics of r and K cells mixed populations.",
        "Black boxes and lines represent simulation results. Gray boxes and lines represent observations.",
        "Supplementary Figure 6 | Detachment curves of r and K cells under trypsinization.",
        "Cells detached under trypsinization were counted every minute.",
        "Supplementary Figure 9 | Growth model fitting.",
        "Exponential (R- Squared number 0.739; p=3.02e-05), Gompertz (R- Squared number 0.828; p=7.7e-05) and Logistic (R- Squared number 0.856; p=4.95e-05).",
        "Supplementary Figure 10 | Carrying capacity estimation.",
        "The functions of density curve of r and K cells were estimated as 228280/(1+83.485 exp(-0.80585 x)) and 239120/(1+728.8 exp(-1.0549 x)), respectively.",
        "Supplementary Figure 11 | Predicted dynamics of r and K cells mixed populations.",
        "The populations were cultured under high density based on the density dependent population growth model.",
        "Supplementary Figure 12 | The dynamics of r and K cell mixture populations.",
        "Each panel shows 100 simulation predictions of a mixture population.",
        "Supplementary Figure 13 | The spatial computational model of population growth.",
        "The cell growth space was assumed to be a two-dimensional planar grid with migration, division, and density dependent regions.",
        "Supplementary Figure 14 | Density-dependent cell size.",
        "Fluorescence imaging of cells at two densities.",
        "Supplementary Figure 15 | Ratio of migrated cells.",
        "The data were collected using a trans-well migration assay.",
        "Supplementary Table 1 | The number of DEGs across comparisons.",
        "Supplementary Table 2 | Enrichment of DEGs in r and K cells under low-density.",
        "Supplementary Table 3 | Top 25 pathways enriched in r and K cells under crowed culture.",
        "Supplementary Table 5 |",
        "Supplementary Table 6 | The abbravations of KEGG patways in Figure 2d",
    ]
    tables = [
        [["Comparisons", "High-expressed genes number", "Low-expressed genes number", "Total DEGs number"], ["KL vs. rL", "1748", "1413", "3161"]],
        [["KEGG Pathway", "Count", "%", "P-Value"], ["Spliceosome", "54", "1.7", "1.00E-10"]],
        [["Term", "Count", "%", "PValue", "Fold Enrichment"], ["Proteasome", "19", "1.073", "9.32E-09", "4.892"]],
        [
            ["Samples", "IN_G", "IN_R", "G3K", "R1K", "G3r", "R1r"],
            ["1", "1.18205451", "0.9099664", "1.05785739", "0.4169925", "1.1529676", "1.36096405"],
            ["2", "0.78450625", "1.09276245", "0.96578718", "0.3", "0.63894989", "1.20447356"],
        ],
        [["Abbravation", "Pathway"], ["AJ", "Adherens junction"]],
    ]
    _write_minimal_docx(root / "Supplementary Figures and Tables revision_2nd.docx", figure_paragraphs, tables)
    _write_minimal_docx(
        root / "Supplementary data- Materials and Methods.docx",
        [
            "Materials and Methods",
            "The correlations are maximized when (, ) = (2.5,0).",
            "The spatial model uses a two-dimensional grid, migration, division, and density-dependent space.",
        ],
        [],
    )
    for name in ("Supplementary_Data_1.txt", "Supplementary_Data_2.txt", "Supplementary_Data_3.txt"):
        (root / name).write_text("Gene\tPPEE\tPPDE\tPostFC\tRealFC\nMOCK\t0\t1\t1.0\t1.0\n")
    (root / "Supplementary_Data_4.vcf").write_text("##fileformat=VCFv4.2\n")
    return root


def _write_minimal_docx(path: Path, paragraphs: list[str], tables: list[list[list[str]]]) -> None:
    w_ns = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
    parts = [f"<w:p><w:r><w:t>{escape(text)}</w:t></w:r></w:p>" for text in paragraphs]
    for table in tables:
        row_xml = []
        for row in table:
            cells = "".join(
                f"<w:tc><w:p><w:r><w:t>{escape(cell)}</w:t></w:r></w:p></w:tc>"
                for cell in row
            )
            row_xml.append(f"<w:tr>{cells}</w:tr>")
        parts.append(f"<w:tbl>{''.join(row_xml)}</w:tbl>")
    document = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        f'<w:document xmlns:w="{w_ns}"><w:body>{"".join(parts)}</w:body></w:document>'
    )
    with zipfile.ZipFile(path, "w") as archive:
        archive.writestr("[Content_Types].xml", "<Types/>")
        archive.writestr("word/document.xml", document)
        archive.writestr("word/media/image1.png", b"mock embedded media placeholder")
